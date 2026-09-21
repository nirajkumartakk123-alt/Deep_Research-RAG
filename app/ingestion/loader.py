"""
Format-specific document loading.

Each loader function takes a file path and returns a list of
PageContent objects — one per page (PDF), section (Markdown/HTML/DOCX
split on headings), or a single entry for formats with no natural
subdivision (TXT).

Known limitations (documented deliberately, not hidden):
- PDF: no cross-page section/heading detection, only page numbers.
- DOCX: no native page numbers exist in the format itself (page breaks
  are a rendering concept, not stored data) - only heading-based sections.
- HTML: heading/section grouping is a simple linear top-down walk, not
  a full DOM-aware structure parse - deeply nested layouts may merge
  sections that a human would consider separate.
These are acceptable trade-offs for an MVP and are good, honest talking
points for an interview about what a "real" ingestion system would add.
"""
import re
from dataclasses import dataclass
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.core.exceptions import DocumentExtractionError, UnsupportedFileTypeError


@dataclass
class PageContent:
    text: str
    page: int | None
    section: str | None


@dataclass
class LoadedDocument:
    pages: list[PageContent]
    document_type: str
    original_filename: str


def _load_txt(path: Path) -> list[PageContent]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [PageContent(text=text, page=None, section=None)]


def _load_markdown(path: Path) -> list[PageContent]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    heading_pattern = re.compile(r"^#{1,6}\s+(.*)")

    sections: list[PageContent] = []
    current_heading: str | None = None
    current_lines: list[str] = []

    for line in text.splitlines():
        match = heading_pattern.match(line)
        if match:
            if current_lines:
                sections.append(
                    PageContent(text="\n".join(current_lines), page=None, section=current_heading)
                )
                current_lines = []
            current_heading = match.group(1).strip()
        else:
            current_lines.append(line)

    if current_lines:
        sections.append(PageContent(text="\n".join(current_lines), page=None, section=current_heading))

    return sections


def _load_html(path: Path) -> list[PageContent]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(raw, "html.parser")

    for tag in soup(["script", "style"]):
        tag.decompose()

    body = soup.body or soup

    sections: list[PageContent] = []
    current_heading: str | None = None
    current_parts: list[str] = []

    for element in body.find_all(["h1", "h2", "h3", "h4", "p", "li"], recursive=True):
        # Skip elements nested inside a <p> or <li> we've already captured,
        # to avoid double-counting text (e.g. a <p> containing inline tags).
        if element.find_parent(["p", "li"]):
            continue

        text = element.get_text(" ", strip=True)
        if not text:
            continue

        if element.name in ("h1", "h2", "h3", "h4"):
            if current_parts:
                sections.append(
                    PageContent(text="\n".join(current_parts), page=None, section=current_heading)
                )
                current_parts = []
            current_heading = text
        else:
            current_parts.append(text)

    if current_parts:
        sections.append(PageContent(text="\n".join(current_parts), page=None, section=current_heading))

    if not sections:
        full_text = soup.get_text(" ", strip=True)
        sections = [PageContent(text=full_text, page=None, section=None)]

    return sections


def _load_docx(path: Path) -> list[PageContent]:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise DocumentExtractionError(f"Could not read DOCX '{path.name}': {exc}") from exc

    sections: list[PageContent] = []
    current_heading: str | None = None
    current_parts: list[str] = []

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = (para.style.name or "").lower()
        is_heading = style_name.startswith("heading") or style_name == "title"

        if is_heading:
            if current_parts:
                sections.append(
                    PageContent(text="\n".join(current_parts), page=None, section=current_heading)
                )
                current_parts = []
            current_heading = text
        else:
            current_parts.append(text)

    if current_parts:
        sections.append(PageContent(text="\n".join(current_parts), page=None, section=current_heading))

    return sections


def _load_pdf(path: Path) -> list[PageContent]:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise DocumentExtractionError(f"Could not read PDF '{path.name}': {exc}") from exc

    pages: list[PageContent] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise DocumentExtractionError(
                f"Failed to extract text from page {i} of '{path.name}': {exc}"
            ) from exc
        pages.append(PageContent(text=text, page=i, section=None))

    return pages


_LOADERS = {
    ".txt": _load_txt,
    ".md": _load_markdown,
    ".html": _load_html,
    ".htm": _load_html,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}

ALLOWED_EXTENSIONS = set(_LOADERS.keys())


def load_document(path: Path, original_filename: str) -> LoadedDocument:
    ext = Path(original_filename).suffix.lower()
    loader_fn = _LOADERS.get(ext)

    if loader_fn is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}' for '{original_filename}'. "
            f"Supported types: {sorted(ALLOWED_EXTENSIONS)}"
        )

    pages = loader_fn(path)
    return LoadedDocument(pages=pages, document_type=ext.lstrip("."), original_filename=original_filename)