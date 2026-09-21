import pytest

from app.core.exceptions import UnsupportedFileTypeError
from app.ingestion.loader import load_document


def test_load_txt(tmp_path):
    p = tmp_path / "sample.txt"
    p.write_text("Hello world.")
    doc = load_document(p, "sample.txt")
    assert doc.pages[0].text.strip() == "Hello world."
    assert doc.document_type == "txt"


def test_load_markdown_sections(tmp_path):
    p = tmp_path / "sample.md"
    p.write_text("# Intro\nHello\n\n## Details\nMore text")
    doc = load_document(p, "sample.md")
    sections = [pg.section for pg in doc.pages]
    assert "Intro" in sections
    assert "Details" in sections


def test_load_html_extracts_text_and_heading(tmp_path):
    p = tmp_path / "sample.html"
    p.write_text("<html><body><h1>Title</h1><p>Some content</p></body></html>")
    doc = load_document(p, "sample.html")
    assert any("Some content" in pg.text for pg in doc.pages)
    assert any(pg.section == "Title" for pg in doc.pages)


def test_load_docx(tmp_path):
    from docx import Document as DocxDocument

    p = tmp_path / "sample.docx"
    d = DocxDocument()
    d.add_heading("Chapter 1", level=1)
    d.add_paragraph("Body text here.")
    d.save(p)

    doc = load_document(p, "sample.docx")
    assert any("Body text here." in pg.text for pg in doc.pages)
    assert any(pg.section == "Chapter 1" for pg in doc.pages)


def test_load_pdf(tmp_path):
    from pypdf import PdfWriter

    p = tmp_path / "sample.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(p, "wb") as f:
        writer.write(f)

    doc = load_document(p, "sample.pdf")
    assert doc.pages[0].page == 1


def test_unsupported_extension_raises(tmp_path):
    p = tmp_path / "sample.xyz"
    p.write_text("data")
    with pytest.raises(UnsupportedFileTypeError):
        load_document(p, "sample.xyz")